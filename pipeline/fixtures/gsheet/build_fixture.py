#!/usr/bin/env python3
"""
build_fixture.py — the SYNTHETIC spreadsheet the importer is tested against.

This is not the real community sheet and nothing in it is a real result
except the one row deliberately mirroring the VOD milestone (Al Qadsiah vs
Twisted Minds, Nepal, 2026-07-04), which exists so reconciliation has a real
VOD map to link to. Every other team, player and score is invented and
named to look invented ('Fixture …', 'Team Nova').

It exercises every layout the importer claims to support, so a test failing
here means a claim stopped being true:

  About          credit text ('Sheet made by …') — no data
  EMEA 2026      one row per map, TWO-ROW header (merged 'Team 1' group
                 label over Ban / Tank / DPS 1 …), merged date cell, a
                 repeated header row, a blank separator, a junk 'Total' row,
                 'Juno / Lúcio' in one cell, McCree (renamed hero), a typo
                 ('Jnuo'), a hyperlink VOD cell
  Korea Stage 2  one row per PLAYER per map (long format), 'ZOX' vs 'zox'
                 casing, text dates ('Jul 12, 2026')
  Bans           one row per hero ban
  Series         one row per series (series score)
  All Maps       the same EMEA map again (an identical duplicate) and a
                 conflicting duplicate (different winner)
  Scratch        free text nothing recognises — preserved, never parsed

Run:  python3 pipeline/fixtures/gsheet/build_fixture.py
Writes pipeline/fixtures/gsheet/owcs_sheet_fixture.xlsx (committed).
"""
from __future__ import annotations

import datetime as dt
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "owcs_sheet_fixture.xlsx")

T1 = ["Tank", "DPS 1", "DPS 2", "Support 1", "Support 2"]


def emea(ws) -> None:
    ws.append(["OWCS 2026 EMEA — fixture tab (synthetic)"])
    top = ["", "", "", "", "", "", "", "", "", "", "Team 1", "", "", "", "", "", "Team 2", "", "", "", "", "", ""]
    sub = ["Date", "Stage", "Team 1", "Team 2", "Map #", "Map", "Mode", "Winner", "Pick", "Score",
           "Ban"] + T1 + ["Ban"] + T1 + ["VOD"]
    ws.append(top)
    ws.append(sub)
    ws.merge_cells("K2:P2")
    ws.merge_cells("Q2:V2")
    rows = [
        # The one row that mirrors the real VOD milestone.
        [dt.datetime(2026, 7, 4), "Stage 2 Playoffs", "Al Qadsiah", "Twisted Minds", 1, "Nepal", "Control",
         "Twisted Minds", "Al Qadsiah", "1-2", "Tracer",
         "Mauga", "Shion", "Symmetra", "Kiriko", "Juno / Lúcio",
         "Ana", "D.Va", "Sojourn", "Sym", "Lucio", "Kiriko",
         "https://www.youtube.com/live/jkSiX___Qwc"],
        [None, "Stage 2 Playoffs", "Al Qadsiah", "Twisted Minds", 2, "Rialto", "Escort",
         "Al Qadsiah", "Twisted Minds", "3-2", "Ana",
         "Winston", "McCree", "Tracer", "Lucio", "Kiriko",
         "Sigma", "D.Va", "Sojourn", "Genji", "Lúcio", "Jnuo",
         None],
        [None, "Stage 2 Playoffs", "Al Qadsiah", "Twisted Minds", 3, "Colosseo", "Push",
         "TM", "decider", "", "Juno",
         "Mauga", "Shion", "Sojourn", "Kiriko", "Baptiste",
         "Winston", "Hazard", "Sojourn", "Shion", "Lúcio", "Ana", None],
    ]
    for r in rows:
        ws.append(r)
    ws.merge_cells("A4:A6")                    # one date over three maps
    ws["W4"].hyperlink = "https://www.youtube.com/live/jkSiX___Qwc?t=1805"
    ws.append([])                              # blank separator
    ws.append(sub)                             # repeated header row
    ws.append([dt.datetime(2026, 7, 5), "Stage 2 Playoffs", "Fixture Falcons", "Team Nova", 1,
               "Kings Row", "Hybrid", "Team 2", "Team 1", "3-2", "Sojourn",
               "Reinhardt", "Cassidy", "Soldier 76", "Ana", "Brigitte",
               "Juno", "Wrecking Ball", "Tracer", "Genji", "Lúcio", "Kiriko", None])
    ws.append([None, "Stage 2 Playoffs", "Fixture Falcons", "Team Nova", 2,
               "Aatlis", "Flashpoint", "Fixture Falcons", "Team Nova", "", "Kiriko",
               "Ramattra", "Junkrat", "Hanzo", "Ana", "Juno",
               "Sigma", "D.Va", "Tracer", "Echo", "Lifeweaver", "Wuyang", None])
    ws.append(["", "", "", "", "", "Total", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""])


def korea(ws) -> None:
    ws.append(["Date", "Event", "Team", "Opponent", "Map #", "Map", "Player", "Role", "Hero 1", "Hero 2", "Hero 3"])
    rows = [
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Raccoon", "Fixture Zeta", 1, "Busan", "Alpha", "Tank", "Winston", "D.Va", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Raccoon", "Fixture Zeta", 1, "Busan", "Bravo", "Damage", "Tracer", "", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Raccoon", "Fixture Zeta", 1, "Busan", "Charlie", "Damage", "Genji", "Sojourn", "Echo"),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Raccoon", "Fixture Zeta", 1, "Busan", "Delta", "Support", "Ana", "", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Raccoon", "Fixture Zeta", 1, "Busan", "Echoo", "Support", "Kiriko", "", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Zeta", "Fixture Raccoon", 1, "Busan", "ZOX", "Support", "Juno", "Lucio", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Zeta", "Fixture Raccoon", 1, "Busan", "Golf", "Tank", "Sigma", "", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Zeta", "Fixture Raccoon", 2, "Midtown", "zox", "Support", "Lucio", "", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Fixture Zeta", "Fixture Raccoon", 2, "Midtown", "Golf", "Tank", "Reinhardt", "", ""),
        ("Jul 12, 2026", "Fixture Cup Korea Stage 2", "Mystery Org", "", 3, "Ilios", "Kilo", "Tank", "Orisa", "", ""),
    ]
    for r in rows:
        ws.append(list(r))


def bans(ws) -> None:
    ws.append(["Date", "Team 1", "Team 2", "Map #", "Team", "Banned Hero"])
    ws.append(["Jul 12, 2026", "Fixture Raccoon", "Fixture Zeta", 1, "Fixture Raccoon", "Juno"])
    ws.append(["Jul 12, 2026", "Fixture Raccoon", "Fixture Zeta", 1, "Fixture Zeta", "Tracer"])
    ws.append(["Jul 12, 2026", "Fixture Raccoon", "Fixture Zeta", 2, "Fixture Zeta", "Sojourn"])


def series(ws) -> None:
    ws.append(["Date", "Event", "Team 1", "Team 2", "Series Score", "Winner"])
    # Consistent with the EMEA tab's map winners (TM, QAD, TM).
    ws.append([dt.datetime(2026, 7, 4), "OWCS 2026 EMEA", "Al Qadsiah", "Twisted Minds", "1-2", "Twisted Minds"])
    # Deliberately INCONSISTENT with its own map rows (1-1): the importer
    # must flag a sheet that disagrees with itself, not pick a side.
    ws.append([dt.datetime(2026, 7, 5), "OWCS 2026 EMEA", "Fixture Falcons", "Team Nova", "2-0", "Fixture Falcons"])
    ws.append([dt.datetime(2026, 7, 20), "OWCS 2026 EMEA", "Fixture Falcons", "Fixture Lions", "3-0", "Fixture Falcons"])


def all_maps(ws) -> None:
    ws.append(["Date", "Team 1", "Team 2", "Map #", "Map", "Winner", "Region"])
    ws.append([dt.datetime(2026, 7, 5), "Fixture Falcons", "Team Nova", 2, "Aatlis", "Fixture Falcons", "EMEA"])
    ws.append([dt.datetime(2026, 7, 5), "Fixture Falcons", "Team Nova", 1, "Kings Row", "Fixture Falcons", "EMEA"])


def build(path: str = OUT) -> str:
    wb = openpyxl.Workbook()
    about = wb.active
    about.title = "About"
    about.append(["OWCS match tracker — SYNTHETIC FIXTURE for tests"])
    about.append(["Sheet made by Fixture Editor (not a real person)"])
    emea(wb.create_sheet("EMEA 2026"))
    korea(wb.create_sheet("Korea Stage 2"))
    bans(wb.create_sheet("Bans"))
    series(wb.create_sheet("Series"))
    all_maps(wb.create_sheet("All Maps"))
    s = wb.create_sheet("Scratch")
    s.append(["random thoughts", "nothing here is a table"])
    s.append(["more text"])
    fixed = dt.datetime(2026, 1, 1)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())
